from __future__ import annotations

import csv
import itertools
import json
import math
from pathlib import Path
from typing import Any

import pandas as pd

from data_agent_baseline.benchmark.schema import PublicTask


_DESCRIBE_DEFAULT_MAX_ROWS = 10_000
_HEAD_DEFAULT_N = 10
# §3.2 size-aware streaming defaults — never load more than this from a
# single file in one preview call. Agents that need more should drive
# `execute_python` with explicit chunked reads.
_DOC_DEFAULT_MAX_BYTES = 256_000          # ~250 KB raw text cap
_JSON_DEFAULT_MAX_BYTES = 512_000         # ~500 KB raw JSON cap
_LARGE_FILE_BYTES = 50 * 1024 * 1024      # 50 MB threshold for catalog hint


def resolve_context_path(task: PublicTask, relative_path: str) -> Path:
    candidate = (task.context_dir / relative_path).resolve()
    context_root = task.context_dir.resolve()
    if context_root not in candidate.parents and candidate != context_root:
        raise ValueError(f"Path escapes context dir: {relative_path}")
    if not candidate.exists():
        raise FileNotFoundError(f"Missing context asset: {relative_path}")
    return candidate


def _ext_kind(path: Path) -> str:
    """Coarse classification of a context file by extension.

    Used by ``list_context_tree`` to surface the hint the agent needs to
    pick the right reader without sampling the file. Anything outside the
    known set falls back to ``"other"``.
    """
    suffix = path.suffix.lower()
    return {
        ".csv": "csv",
        ".tsv": "csv",
        ".json": "json",
        ".db": "sqlite",
        ".sqlite": "sqlite",
        ".sqlite3": "sqlite",
        ".md": "doc",
        ".txt": "doc",
        ".pdf": "pdf",
        ".xlsx": "excel",
        ".xlsm": "excel",
        ".parquet": "parquet",
        ".png": "image",
        ".jpg": "image",
        ".jpeg": "image",
        ".gif": "image",
        ".bmp": "image",
        ".webp": "image",
        ".tiff": "image",
        ".tif": "image",
        ".zip": "archive",
        ".tar": "archive",
        ".tgz": "archive",
        ".gz": "archive",
    }.get(suffix, "other")


def list_context_tree(task: PublicTask, *, max_depth: int = 4) -> dict[str, object]:
    """Catalog of every file/dir under context/ with size + ext-based kind.

    §3.2 hint: when a file exceeds ``_LARGE_FILE_BYTES`` (50 MB) the entry
    flips ``large=True`` so the agent prefers the streaming-friendly
    readers (dataframe_describe, dataframe_head, read_parquet, etc.)
    instead of the whole-file ones (read_csv, read_json, read_doc).
    """
    entries: list[dict[str, object]] = []

    def walk(path: Path, depth: int) -> None:
        if depth > max_depth:
            return
        for child in sorted(path.iterdir(), key=lambda item: (item.is_file(), item.name)):
            rel_path = child.relative_to(task.context_dir).as_posix()
            if child.is_dir():
                entries.append({"path": rel_path, "kind": "dir", "size": None, "ext_kind": "dir"})
                walk(child, depth + 1)
                continue
            size = child.stat().st_size
            entry: dict[str, object] = {
                "path": rel_path,
                "kind": "file",
                "size": size,
                "ext_kind": _ext_kind(child),
            }
            if size >= _LARGE_FILE_BYTES:
                entry["large"] = True
            entries.append(entry)

    walk(task.context_dir, 1)
    return {
        "root": str(task.context_dir),
        "entries": entries,
        "large_file_threshold_bytes": _LARGE_FILE_BYTES,
    }


def read_csv_preview(task: PublicTask, relative_path: str, *, max_rows: int = 20) -> dict[str, object]:
    """Header + first N rows of a CSV. Streams via ``csv.reader``: never
    loads more than ``max_rows + 1`` rows into memory regardless of file
    size. ``row_count_known`` distinguishes the cheap path (we read every
    row) from the bounded path (we capped at ``max_rows``)."""
    path = resolve_context_path(task, relative_path)
    n = max(0, int(max_rows))
    with path.open(newline="") as handle:
        reader = csv.reader(handle)
        try:
            header = next(reader)
        except StopIteration:
            return {
                "path": relative_path,
                "columns": [],
                "rows": [],
                "row_count": 0,
                "row_count_known": True,
            }
        sliced = list(itertools.islice(reader, n + 1))
        truncated = len(sliced) > n
        data_rows = sliced[:n]
    return {
        "path": relative_path,
        "columns": header,
        "rows": data_rows,
        "row_count": len(data_rows) if not truncated else None,
        "row_count_known": not truncated,
        "truncated": truncated,
    }


def read_json_preview(
    task: PublicTask,
    relative_path: str,
    *,
    max_chars: int = 4000,
    max_bytes: int = _JSON_DEFAULT_MAX_BYTES,
) -> dict[str, object]:
    """JSON pretty-print preview with size guard.

    Files over ``max_bytes`` are NOT loaded — we return a metadata-only
    response so the agent picks ``execute_python`` for streaming parsing
    (e.g. ``ijson``) instead. Files within budget are decoded, pretty-
    printed, and truncated to ``max_chars`` of preview text.
    """
    path = resolve_context_path(task, relative_path)
    size = path.stat().st_size
    if size > max_bytes:
        return {
            "path": relative_path,
            "size_bytes": size,
            "max_bytes": max_bytes,
            "skipped": True,
            "reason": (
                f"file is {size} bytes which exceeds max_bytes={max_bytes}; "
                "use execute_python with streaming JSON parsing (e.g. ijson) "
                "or pass a larger max_bytes"
            ),
        }
    payload = json.loads(path.read_text())
    preview = json.dumps(payload, ensure_ascii=False, indent=2)
    return {
        "path": relative_path,
        "size_bytes": size,
        "preview": preview[:max_chars],
        "truncated": len(preview) > max_chars,
    }


def read_doc_preview(
    task: PublicTask,
    relative_path: str,
    *,
    max_chars: int = 4000,
    max_bytes: int = _DOC_DEFAULT_MAX_BYTES,
) -> dict[str, object]:
    """Text-document preview with byte-bounded read.

    Reads at most ``max_bytes`` from disk (so a 1 GB log file is safe),
    then truncates the decoded text to ``max_chars``. ``read_truncated``
    flips True when the file is bigger than the byte budget — the rest
    of the file is intentionally not on disk-read in this call.
    """
    path = resolve_context_path(task, relative_path)
    size = path.stat().st_size
    cap_bytes = max(0, int(max_bytes))
    with path.open("rb") as handle:
        raw = handle.read(cap_bytes)
    text = raw.decode("utf-8", errors="replace")
    return {
        "path": relative_path,
        "size_bytes": size,
        "preview": text[:max_chars],
        "truncated": len(text) > max_chars,
        "read_truncated": size > cap_bytes,
        "max_bytes": cap_bytes,
    }


def _scrub_for_json(value: Any) -> Any:
    """Convert pandas/numpy scalars to JSON-friendly Python primitives.

    Handles NaN/Inf (which break json.dumps with allow_nan=False downstream)
    and numpy ints/floats. Strings, bools, and None pass through untouched.
    """
    if value is None:
        return None
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return value
    if isinstance(value, (int, str, bool)):
        return value
    # numpy scalar fallback
    item = getattr(value, "item", None)
    if callable(item):
        try:
            scalar = item()
        except Exception:  # noqa: BLE001
            return str(value)
        if isinstance(scalar, float) and (math.isnan(scalar) or math.isinf(scalar)):
            return None
        return scalar
    return str(value)


def dataframe_describe(
    task: PublicTask,
    relative_path: str,
    *,
    max_rows: int = _DESCRIBE_DEFAULT_MAX_ROWS,
    head_rows: int = 5,
) -> dict[str, object]:
    """Compressed metadata for a tabular file (CSV).

    Reads up to `max_rows` rows from the file and returns dtypes,
    pandas.describe() summary, and a small head preview. Designed as a
    cheap prepass for large CSVs: the agent gets shape/dtype/stat hints
    without us shipping hundreds of MB of raw rows back through the
    observation channel.
    """
    path = resolve_context_path(task, relative_path)
    suffix = path.suffix.lower()
    if suffix not in {".csv", ".tsv"}:
        raise ValueError(
            f"dataframe_describe only supports .csv or .tsv files, got {suffix!r}."
        )
    sep = "\t" if suffix == ".tsv" else ","
    capped_rows = max(1, int(max_rows))

    try:
        df = pd.read_csv(path, sep=sep, nrows=capped_rows, low_memory=False)
    except Exception as exc:  # noqa: BLE001
        return {"path": relative_path, "error": f"pandas.read_csv failed: {exc}"}

    truncated = len(df) >= capped_rows  # we cannot tell exactly without rereading

    describe_df = df.describe(include="all").transpose()
    describe: dict[str, dict[str, Any]] = {}
    for column_name, row in describe_df.iterrows():
        describe[str(column_name)] = {
            stat: _scrub_for_json(value) for stat, value in row.items()
        }

    head_df = df.head(max(0, int(head_rows)))
    head_rows_payload = [
        [_scrub_for_json(cell) for cell in record]
        for record in head_df.to_records(index=False).tolist()
    ]

    return {
        "path": relative_path,
        "shape": [int(df.shape[0]), int(df.shape[1])],
        "rows_read_capped_at": capped_rows,
        "rows_may_be_truncated": truncated,
        "columns": [str(name) for name in df.columns],
        "dtypes": {str(name): str(dtype) for name, dtype in df.dtypes.items()},
        "describe": describe,
        "head_rows": head_rows_payload,
    }


def dataframe_head(
    task: PublicTask,
    relative_path: str,
    *,
    n: int = _HEAD_DEFAULT_N,
) -> dict[str, object]:
    """Read just the first `n` rows of a CSV/TSV with proper dtype inference.

    Lighter than dataframe_describe: does not run describe(). Useful for
    confirming column names and a few sample values cheaply.
    """
    path = resolve_context_path(task, relative_path)
    suffix = path.suffix.lower()
    if suffix not in {".csv", ".tsv"}:
        raise ValueError(
            f"dataframe_head only supports .csv or .tsv files, got {suffix!r}."
        )
    sep = "\t" if suffix == ".tsv" else ","
    n_int = max(1, int(n))

    try:
        df = pd.read_csv(path, sep=sep, nrows=n_int, low_memory=False)
    except Exception as exc:  # noqa: BLE001
        return {"path": relative_path, "error": f"pandas.read_csv failed: {exc}"}

    rows_payload = [
        [_scrub_for_json(cell) for cell in record]
        for record in df.to_records(index=False).tolist()
    ]
    return {
        "path": relative_path,
        "shape": [int(df.shape[0]), int(df.shape[1])],
        "columns": [str(name) for name in df.columns],
        "dtypes": {str(name): str(dtype) for name, dtype in df.dtypes.items()},
        "rows": rows_payload,
    }


# ─────────────────────────────────────────────────────────────────────────────
# Format dispatcher (§3.1) — readers for the formats the public 50-task set
# does not contain but the hidden Phase-1 / Phase-2 set may. All deterministic;
# no auxiliary LLM, no vision/OCR pipeline, no network calls (single-model
# policy + rules §runtime network ban).
# ─────────────────────────────────────────────────────────────────────────────


_PDF_DEFAULT_MAX_PAGES = 5
_PDF_DEFAULT_MAX_CHARS = 4000
_EXCEL_DEFAULT_MAX_ROWS = 20
_PARQUET_DEFAULT_MAX_ROWS = 20
_ARCHIVE_DEFAULT_MAX_ENTRIES = 200
_IMAGE_SUFFIXES: frozenset[str] = frozenset(
    {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff", ".tif"}
)


def read_pdf_preview(
    task: PublicTask,
    relative_path: str,
    *,
    max_pages: int = _PDF_DEFAULT_MAX_PAGES,
    max_chars: int = _PDF_DEFAULT_MAX_CHARS,
) -> dict[str, object]:
    """Per-page text extract from a .pdf using pypdf (deterministic, no OCR).

    Caps both the number of pages walked AND the total characters returned
    so a 200-page PDF cannot blow up the observation budget.
    """
    from pypdf import PdfReader

    path = resolve_context_path(task, relative_path)
    if path.suffix.lower() != ".pdf":
        raise ValueError(
            f"read_pdf_preview only supports .pdf files, got {path.suffix!r}."
        )
    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001
        return {"path": relative_path, "error": f"pypdf failed: {exc}"}

    page_count = len(reader.pages)
    pages_to_read = min(page_count, max(1, int(max_pages)))
    char_budget = max(0, int(max_chars))
    pages_out: list[dict[str, object]] = []

    for idx in range(pages_to_read):
        try:
            text = (reader.pages[idx].extract_text() or "").strip()
        except Exception as exc:  # noqa: BLE001
            text = f"[page {idx + 1} extract failed: {exc}]"
        if char_budget <= 0:
            pages_out.append(
                {"page": idx + 1, "preview": "[budget exhausted]", "char_count": len(text)}
            )
            break
        snippet = text[:char_budget]
        if len(text) > char_budget:
            snippet = snippet.rstrip() + "\n... [page truncated]"
            char_budget = 0
        else:
            char_budget -= len(snippet)
        pages_out.append({"page": idx + 1, "preview": snippet, "char_count": len(text)})

    return {
        "path": relative_path,
        "page_count": page_count,
        "pages_returned": len(pages_out),
        "pages": pages_out,
        "max_pages": pages_to_read,
        "max_chars": int(max_chars),
    }


def read_excel_preview(
    task: PublicTask,
    relative_path: str,
    *,
    sheet: str | None = None,
    max_rows: int = _EXCEL_DEFAULT_MAX_ROWS,
) -> dict[str, object]:
    """Read sheet names + sampled rows from .xlsx/.xlsm using openpyxl.

    Defaults to the first sheet when ``sheet`` is None. The first row is
    treated as the header. Read-only mode keeps memory bounded.
    """
    from openpyxl import load_workbook

    path = resolve_context_path(task, relative_path)
    suffix = path.suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError(
            f"read_excel_preview supports .xlsx/.xlsm, got {suffix!r}."
        )
    try:
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return {"path": relative_path, "error": f"openpyxl failed: {exc}"}

    sheets = list(wb.sheetnames)
    target_sheet = sheet if (sheet and sheet in sheets) else (sheets[0] if sheets else None)
    if target_sheet is None:
        wb.close()
        return {"path": relative_path, "sheets": [], "error": "workbook has no sheets"}

    ws = wb[target_sheet]
    n_int = max(1, int(max_rows))
    collected: list[list[object]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > n_int:  # +1 for header row
            break
        collected.append([_scrub_for_json(v) for v in row])
    wb.close()

    header = collected[0] if collected else []
    body = collected[1:] if len(collected) > 1 else []
    return {
        "path": relative_path,
        "sheets": sheets,
        "sheet_used": target_sheet,
        "header": header,
        "rows": body,
        "rows_returned": len(body),
        "max_rows": n_int,
    }


def read_parquet_preview(
    task: PublicTask,
    relative_path: str,
    *,
    max_rows: int = _PARQUET_DEFAULT_MAX_ROWS,
) -> dict[str, object]:
    """Parquet schema + sampled head using pyarrow."""
    import pyarrow.parquet as pq

    path = resolve_context_path(task, relative_path)
    if path.suffix.lower() != ".parquet":
        raise ValueError(
            f"read_parquet_preview only supports .parquet, got {path.suffix!r}."
        )
    try:
        pf = pq.ParquetFile(str(path))
    except Exception as exc:  # noqa: BLE001
        return {"path": relative_path, "error": f"pyarrow failed: {exc}"}

    schema = {field.name: str(field.type) for field in pf.schema_arrow}
    columns = list(schema.keys())
    total_rows = int(pf.metadata.num_rows)
    n_int = max(1, int(max_rows))
    rows: list[list[object]] = []
    if pf.num_row_groups > 0:
        head_table = pf.read_row_group(0).slice(0, n_int)
        for record in head_table.to_pylist():
            rows.append([_scrub_for_json(record.get(c)) for c in columns])

    return {
        "path": relative_path,
        "row_count": total_rows,
        "row_group_count": pf.num_row_groups,
        "schema": schema,
        "columns": columns,
        "rows": rows,
        "max_rows": n_int,
    }


def read_image_meta(
    task: PublicTask,
    relative_path: str,
) -> dict[str, object]:
    """Image metadata only — format, mode, dimensions, byte size.

    ``team1438`` single-model policy: we do NOT run OCR or any vision LLM
    pipeline on the file. The agent only sees structural metadata so it
    can decide whether a vision-bearing answer is even feasible (likely
    not, given the policy).
    """
    from PIL import Image

    path = resolve_context_path(task, relative_path)
    suffix = path.suffix.lower()
    if suffix not in _IMAGE_SUFFIXES:
        raise ValueError(
            f"read_image_meta supports common raster formats, got {suffix!r}."
        )
    try:
        with Image.open(str(path)) as img:
            return {
                "path": relative_path,
                "format": img.format,
                "mode": img.mode,
                "size_pixels": [int(img.size[0]), int(img.size[1])],
                "size_bytes": int(path.stat().st_size),
                "policy_note": (
                    "Single-model policy: only metadata is exposed; no "
                    "vision/OCR pipeline runs on this file."
                ),
            }
    except Exception as exc:  # noqa: BLE001
        return {"path": relative_path, "error": f"PIL failed: {exc}"}


def inspect_file(task: PublicTask, relative_path: str) -> dict[str, object]:
    """Hierarchical catalog: dispatch by extension to extract per-file
    metadata in one cheap call so the agent does not have to guess which
    reader to invoke first.

    Returns a uniform envelope with the inferred ``ext_kind`` plus
    type-specific summary fields:

    - csv/tsv  → ``columns``, sample shape from a 5-row read
    - json     → ``top_level_type`` + (object) ``keys`` or (array) ``length`` + sample
    - sqlite   → ``tables``, columns per table (via inspect_sqlite_schema)
    - pdf      → ``page_count`` (no text extracted in this call)
    - xlsx     → ``sheets`` + first sheet header
    - parquet  → ``schema``, ``row_count``, ``row_group_count``
    - image    → metadata via read_image_meta
    - archive  → entry count + first names via read_archive_listing
    - doc/md/txt → ``size_bytes`` + ``preview`` of first ~500 chars
    - other    → ``size_bytes`` + ``hint`` "use execute_python"

    Always populates ``recommended_tool`` so the agent can pick the next
    step without further inference.
    """
    path = resolve_context_path(task, relative_path)
    if path.is_dir():
        raise ValueError(f"inspect_file expects a file path, got directory: {relative_path}")

    ext_kind = _ext_kind(path)
    size = path.stat().st_size
    base: dict[str, object] = {
        "path": relative_path,
        "ext_kind": ext_kind,
        "size_bytes": size,
        "large": size >= _LARGE_FILE_BYTES,
    }

    if ext_kind == "csv":
        try:
            sep = "\t" if path.suffix.lower() == ".tsv" else ","
            df_head = pd.read_csv(path, sep=sep, nrows=5, low_memory=False)
            base.update(
                {
                    "columns": [str(c) for c in df_head.columns],
                    "head_shape": [int(df_head.shape[0]), int(df_head.shape[1])],
                    "recommended_tool": (
                        "dataframe_describe" if size >= _LARGE_FILE_BYTES else "read_csv"
                    ),
                }
            )
        except Exception as exc:  # noqa: BLE001
            base["error"] = f"pandas.read_csv failed: {exc}"
            base["recommended_tool"] = "read_csv"
        return base

    if ext_kind == "json":
        if size > _JSON_DEFAULT_MAX_BYTES:
            base.update(
                {
                    "skipped": True,
                    "reason": f"file is {size} bytes, exceeds JSON inspect budget",
                    "recommended_tool": "execute_python",
                }
            )
            return base
        try:
            payload = json.loads(path.read_text())
        except Exception as exc:  # noqa: BLE001
            base["error"] = f"json.loads failed: {exc}"
            base["recommended_tool"] = "read_json"
            return base
        if isinstance(payload, dict):
            base.update(
                {
                    "top_level_type": "object",
                    "keys": sorted([str(k) for k in payload.keys()])[:50],
                    "key_count": len(payload),
                }
            )
        elif isinstance(payload, list):
            base.update(
                {
                    "top_level_type": "array",
                    "length": len(payload),
                    "sample_item_type": (
                        type(payload[0]).__name__ if payload else None
                    ),
                }
            )
        else:
            base["top_level_type"] = type(payload).__name__
        base["recommended_tool"] = "read_json"
        return base

    if ext_kind == "sqlite":
        from data_agent_baseline.tools.sqlite import inspect_sqlite_schema

        try:
            schema = inspect_sqlite_schema(path)
            base.update(
                {
                    "tables": [
                        t.get("name") if isinstance(t, dict) else str(t)
                        for t in (schema.get("tables", []) or [])
                    ],
                    "schema": schema,
                    "recommended_tool": "execute_context_sql",
                }
            )
        except Exception as exc:  # noqa: BLE001
            base["error"] = f"sqlite inspection failed: {exc}"
            base["recommended_tool"] = "execute_context_sql"
        return base

    if ext_kind == "pdf":
        try:
            from pypdf import PdfReader

            base["page_count"] = len(PdfReader(str(path)).pages)
        except Exception as exc:  # noqa: BLE001
            base["error"] = f"pypdf inspection failed: {exc}"
        base["recommended_tool"] = "read_pdf"
        return base

    if ext_kind == "excel":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(filename=str(path), read_only=True, data_only=True)
            base["sheets"] = list(wb.sheetnames)
            wb.close()
        except Exception as exc:  # noqa: BLE001
            base["error"] = f"openpyxl inspection failed: {exc}"
        base["recommended_tool"] = "read_excel"
        return base

    if ext_kind == "parquet":
        try:
            import pyarrow.parquet as pq

            pf = pq.ParquetFile(str(path))
            base.update(
                {
                    "row_count": int(pf.metadata.num_rows),
                    "row_group_count": pf.num_row_groups,
                    "schema": {
                        field.name: str(field.type) for field in pf.schema_arrow
                    },
                }
            )
        except Exception as exc:  # noqa: BLE001
            base["error"] = f"pyarrow inspection failed: {exc}"
        base["recommended_tool"] = "read_parquet"
        return base

    if ext_kind == "image":
        try:
            from PIL import Image

            with Image.open(str(path)) as img:
                base.update(
                    {
                        "format": img.format,
                        "mode": img.mode,
                        "size_pixels": [int(img.size[0]), int(img.size[1])],
                    }
                )
        except Exception as exc:  # noqa: BLE001
            base["error"] = f"PIL inspection failed: {exc}"
        base["policy_note"] = (
            "Single-model policy: vision/OCR not invoked. Image content "
            "is not retrievable through the tool surface."
        )
        base["recommended_tool"] = "read_image_meta"
        return base

    if ext_kind == "archive":
        try:
            listing = read_archive_listing(task, relative_path, max_entries=20)
            base["entries_returned"] = listing.get("entries_returned")
            base["archive_type"] = listing.get("archive_type")
            base["sample_entries"] = [
                e.get("name") for e in (listing.get("entries") or [])
            ][:20]
        except Exception as exc:  # noqa: BLE001
            base["error"] = f"archive inspection failed: {exc}"
        base["recommended_tool"] = "read_archive_listing"
        return base

    if ext_kind == "doc":
        try:
            with path.open("rb") as handle:
                raw = handle.read(min(size, 500))
            base["preview"] = raw.decode("utf-8", errors="replace")
        except Exception as exc:  # noqa: BLE001
            base["error"] = f"doc preview failed: {exc}"
        base["recommended_tool"] = "read_doc"
        return base

    # ext_kind == "other"
    base.update(
        {
            "hint": (
                "Unrecognized extension; use execute_python to read raw "
                "bytes or treat as text via read_doc if encoding is utf-8."
            ),
            "recommended_tool": "execute_python",
        }
    )
    return base


def read_archive_listing(
    task: PublicTask,
    relative_path: str,
    *,
    max_entries: int = _ARCHIVE_DEFAULT_MAX_ENTRIES,
) -> dict[str, object]:
    """List entries inside .zip / .tar / .tar.gz / .tgz — names + sizes only.

    Does NOT extract contents. The agent should use this to inventory an
    archive cheaply, then decide whether to read individual files via the
    other readers (likely after extracting separately via execute_python).
    """
    import tarfile
    import zipfile

    path = resolve_context_path(task, relative_path)
    suffix = path.suffix.lower()
    n_int = max(1, int(max_entries))
    entries: list[dict[str, object]] = []
    archive_type: str | None = None

    if suffix == ".zip" or zipfile.is_zipfile(str(path)):
        archive_type = "zip"
        try:
            with zipfile.ZipFile(str(path)) as zf:
                for info in zf.infolist()[:n_int]:
                    entries.append(
                        {
                            "name": info.filename,
                            "size": int(info.file_size),
                            "compressed": int(info.compress_size),
                            "is_dir": info.is_dir(),
                        }
                    )
        except Exception as exc:  # noqa: BLE001
            return {"path": relative_path, "error": f"zipfile failed: {exc}"}
    elif tarfile.is_tarfile(str(path)):
        archive_type = "tar"
        try:
            with tarfile.open(str(path)) as tf:
                for i, member in enumerate(tf):
                    if i >= n_int:
                        break
                    entries.append(
                        {
                            "name": member.name,
                            "size": int(member.size),
                            "is_dir": member.isdir(),
                        }
                    )
        except Exception as exc:  # noqa: BLE001
            return {"path": relative_path, "error": f"tarfile failed: {exc}"}
    else:
        raise ValueError(
            f"read_archive_listing supports .zip/.tar/.tar.gz/.tgz, got {suffix!r}."
        )

    return {
        "path": relative_path,
        "archive_type": archive_type,
        "entries": entries,
        "entries_returned": len(entries),
        "max_entries": n_int,
    }
